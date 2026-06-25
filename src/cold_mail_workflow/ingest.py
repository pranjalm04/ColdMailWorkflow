import csv
import logging
import re
from pathlib import Path

from openpyxl import load_workbook

from .config import CONTACTS_PATH, DEFAULT_ROLE
from .models import Contact

logger = logging.getLogger(__name__)

ALIASES = {
    "company": ["company", "company name", "company name for emails", "organization", "account"],
    "role": ["role", "position", "target role", "job title"],
    "recruiter_email": ["recruiter email", "email", "primary email", "work email"],
    "all_emails": ["all emails", "emails", "other emails", "secondary emails"],
    "recruiter_name": ["recruiter name", "name", "contact name", "full name"],
    "first_name": ["first name", "first"],
    "last_name": ["last name", "last"],
    "title": ["title", "contact title", "recruiter title"],
    "job_url": ["job url", "posting", "job link"],
    "notes": ["notes", "note"],
    "industry": ["industry"],
}

CSV_SUFFIXES = {".csv"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}


_NAME_TOKEN = re.compile(r"^[A-Za-z][A-Za-z.'-]*$")
_SEP = re.compile(r"[\s_\-]+")
_EMAIL_SPLIT = re.compile(r"[;,/|]+")


def _norm(value) -> str:
    return str(value).strip() if value is not None else ""


def _canon(value) -> str:
    return _SEP.sub(" ", str(value).strip().lower()).strip() if value is not None else ""


def _split_emails(raw: str) -> tuple[str, ...]:
    seen: set[str] = set()
    out: list[str] = []
    for part in _EMAIL_SPLIT.split(raw or ""):
        addr = part.strip()
        key = addr.lower()
        if "@" in addr and key not in seen:
            seen.add(key)
            out.append(addr)
    return tuple(out)


def _clean_name(name: str) -> str:
    tokens = [t for t in name.split() if _NAME_TOKEN.match(t) and "{at}" not in t.lower()]
    return " ".join(tokens)


def _read_rows(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix in CSV_SUFFIXES:
        with path.open("r", newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            return [{_canon(k): _norm(v) for k, v in row.items()} for row in reader]
    if suffix in XLSX_SUFFIXES:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            header = [_canon(c) for c in next(rows)]
        except StopIteration:
            wb.close()
            return []
        out = [
            {header[i]: _norm(v) for i, v in enumerate(row) if i < len(header) and header[i]}
            for row in rows
        ]
        wb.close()
        return out
    raise ValueError(f"Unsupported contacts file type '{suffix}' (expected .csv or .xlsx): {path}")


def _pick(row: dict[str, str], field: str) -> str:
    for alias in ALIASES[field]:
        if row.get(alias):
            return row[alias]
    return ""


def _row_to_contact(row: dict[str, str], default_role: str) -> Contact | None:
    company = _pick(row, "company")
    all_emails = _split_emails(_pick(row, "all_emails"))
    email = _pick(row, "recruiter_email") or (all_emails[0] if all_emails else "")
    if not company or not email:
        return None

    name = _pick(row, "recruiter_name")
    if not name:
        name = " ".join(p for p in (_pick(row, "first_name"), _pick(row, "last_name")) if p).strip()
    name = _clean_name(name)

    role = _pick(row, "role") or default_role
    notes = _pick(row, "notes") or _pick(row, "industry")

    return Contact(
        company=company,
        role=role,
        recruiter_email=email,
        recruiter_name=name,
        title=_pick(row, "title"),
        job_url=_pick(row, "job_url"),
        notes=notes,
        all_emails=all_emails,
    )


SUPPORTED_SUFFIXES = CSV_SUFFIXES | XLSX_SUFFIXES


def _read_file_contacts(path: Path, default_role: str) -> list[Contact]:
    rows = _read_rows(path)
    contacts: list[Contact] = []
    skipped = 0
    for line_no, row in enumerate(rows, start=2):
        contact = _row_to_contact(row, default_role)
        if contact is None:
            skipped += 1
            logger.warning("Skipping %s row %d: missing company or recruiter email", path.name, line_no)
            continue
        contacts.append(contact)

    logger.info("Loaded %d contact(s) from %s (skipped %d)", len(contacts), path.name, skipped)
    return contacts


def _contact_files(directory: Path) -> list[Path]:
    return sorted(
        p
        for p in directory.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_SUFFIXES and not p.name.startswith("~$")
    )


def read_contacts(path: Path = CONTACTS_PATH, default_role: str = DEFAULT_ROLE) -> list[Contact]:
    if not path.exists():
        raise FileNotFoundError(f"Contacts path not found: {path}")

    if path.is_dir():
        files = _contact_files(path)
        if not files:
            logger.warning("No .csv/.xlsx contact files found in %s", path)
            return []
        contacts: list[Contact] = []
        for f in files:
            contacts.extend(_read_file_contacts(f, default_role))
        logger.info("Loaded %d contact(s) total from %d file(s) in %s", len(contacts), len(files), path)
        return contacts

    return _read_file_contacts(path, default_role)
